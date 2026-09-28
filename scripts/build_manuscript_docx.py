"""Build the evidence-bound Markdown manuscript using the bundled artifact Python.

Design: narrative_proposal preset, memo_masthead opening. Named academic overrides:
black headings/title, 22 pt compact title, 9 pt table text, no decorative rules.
This renderer does not calculate or certify scientific results.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime,timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor


def property_node(parent,name,attrs):
    existing=parent.find(qn(name))
    if existing is None:
        existing=OxmlElement(name); parent.append(existing)
    for key,value in attrs.items():
        existing.set(qn(key),str(value))
    return existing


def inline(paragraph,text):
    # Preserve readable citation URLs, without renderer-specific citation tokens.
    text=re.sub(r"\[([^\]]+)\]\(([^)]+)\)",r"\1 (\2)",text)
    for piece in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)",text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            paragraph.add_run(piece[2:-2]).bold=True
        elif piece.startswith("*") and piece.endswith("*"):
            paragraph.add_run(piece[1:-1]).italic=True
        elif piece.startswith("`") and piece.endswith("`"):
            run=paragraph.add_run(piece[1:-1]); run.font.name="Consolas"; run.font.size=Pt(9)
        else:
            paragraph.add_run(piece)


def style(doc,name,size,before,after,bold=False,color="000000",line=1.333):
    s=doc.styles[name]; s.font.name="Calibri"; s.font.size=Pt(size)
    s.font.bold=bold; s.font.color.rgb=RGBColor.from_string(color)
    p=s.paragraph_format; p.space_before=Pt(before); p.space_after=Pt(after)
    p.line_spacing=line; p.widow_control=True
    property_node(s.element.get_or_add_rPr(),"w:rFonts",{"w:ascii":"Calibri","w:hAnsi":"Calibri"})
    return s


def add_table(doc,rows):
    n=len(rows[0])
    widths={4:[1600,4800,1480,1480],5:[1700,2100,1760,1900,1900],
            7:[1500,1260,800,980,980,2460,1380]}.get(n)
    if rows[0][0]=="Service" and n==4:
        widths=[1980,1760,2740,2880]
    if n==7 and "Warnings" in rows[0]:
        widths=[1660,1440,1150,1000,1320,1000,1790]
    if widths is None:
        widths=[9360//n]*n; widths[-1]+=9360-sum(widths)
    table=doc.add_table(rows=0,cols=n); table.autofit=False
    props=table._tbl.tblPr
    property_node(props,"w:tblW",{"w:w":9360,"w:type":"dxa"})
    property_node(props,"w:tblInd",{"w:w":120,"w:type":"dxa"})
    property_node(props,"w:tblLayout",{"w:type":"fixed"})
    margins=property_node(props,"w:tblCellMar",{})
    for side,value in (("top",80),("bottom",80),("start",120),("end",120)):
        property_node(margins,"w:"+side,{"w:w":value,"w:type":"dxa"})
    borders=property_node(props,"w:tblBorders",{})
    for side in ("top","left","bottom","right","insideH","insideV"):
        property_node(borders,"w:"+side,{"w:val":"single","w:sz":4,"w:color":"D0D0D0"})
    for col,width in zip(table._tbl.tblGrid.gridCol_lst,widths):
        col.set(qn("w:w"),str(width))
    for i,values in enumerate(rows):
        if len(values)!=n:
            raise ValueError("Unequal table row length")
        row=table.add_row()
        property_node(row._tr.get_or_add_trPr(),"w:cantSplit",{})
        if i==0:
            property_node(row._tr.get_or_add_trPr(),"w:tblHeader",{})
        for j,(cell,text,width) in enumerate(zip(row.cells,values,widths)):
            property_node(cell._tc.get_or_add_tcPr(),"w:tcW",{"w:w":width,"w:type":"dxa"})
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                property_node(cell._tc.get_or_add_tcPr(),"w:shd",{"w:fill":"F4F6F9"})
            p=cell.paragraphs[0]; p.style="Table Text"
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j<2 else WD_ALIGN_PARAGRAPH.CENTER
            inline(p,text)
            if i==0:
                for run in p.runs:
                    run.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(4)


def build(source: Path,output: Path):
    doc=Document()
    for paragraph_style in doc.styles:
        for border in list(paragraph_style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    sec.header_distance=sec.footer_distance=Inches(.492)
    style(doc,"Normal",11,0,8).paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    style(doc,"Title",22,0,10,True,line=1.08)
    style(doc,"Subtitle",11,0,8,color="555555",line=1.15)
    for name,size,before,after in (("Heading 1",16,18,10),("Heading 2",13,12,6),("Heading 3",12,8,4)):
        style(doc,name,size,before,after,True,line=1.15).paragraph_format.keep_with_next=True
    for name in ("Table Text","Review Notice","Table Caption"):
        if name not in doc.styles:
            doc.styles.add_style(name,1)
    style(doc,"Table Text",9,0,2,line=1.08)
    style(doc,"Review Notice",10,4,10,color="555555",line=1.15)
    style(doc,"Table Caption",10,4,4,line=1.15).paragraph_format.keep_with_next=True
    style(doc,"Caption",10,4,8,line=1.15)
    style(doc,"Header",9,0,0,color="555555",line=1)
    style(doc,"Footer",9,0,0,color="555555",line=1)
    sec.header.paragraphs[0].text="Author attribution after LLM rewriting | Review draft"
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run("Review draft | ")
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); p._p.append(field)
    lines=source.read_text(encoding="utf-8").splitlines()
    paragraph=[]; i=0
    def flush():
        if paragraph:
            text=" ".join(paragraph); paragraph.clear()
            role="Table Caption" if re.match(r"Table [A-Z]?\d+\.",text) else "Normal"
            inline(doc.add_paragraph(style=role),text)
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line:
            flush(); continue
        if line.startswith("#"):
            flush(); match=re.match(r"(#+)\s+(.*)",line)
            if not match:
                raise ValueError("Malformed heading")
            depth=len(match[1]); p=doc.add_paragraph(style="Title" if depth==1 else f"Heading {min(depth-1,3)}")
            inline(p,match[2]); continue
        if line.startswith("> "):
            flush(); inline(doc.add_paragraph(style="Review Notice"),line[2:]); continue
        if line.startswith("|"):
            flush(); table_lines=[line]
            while i<len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip()); i+=1
            rows=[[x.strip() for x in r.strip("|").split("|")] for r in table_lines]
            rows=[r for r in rows if not all(re.fullmatch(r":?-+:?",x) for x in r)]
            add_table(doc,rows); continue
        match=re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)",line)
        if match:
            flush(); p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
            run=p.add_run(); run.add_picture(str(source.parent / match[2]),width=Inches(6.5))
            for element in run._r.iter():
                if element.tag.endswith("}docPr"):
                    element.set("descr",match[1])
            inline(doc.add_paragraph(style="Caption"),match[1]); continue
        paragraph.append(line)
    flush()
    doc.core_properties.title=next(line[2:] for line in lines if line.startswith("# "))
    doc.core_properties.author=""; doc.core_properties.last_modified_by=""
    doc.core_properties.created=datetime.now(timezone.utc)
    doc.core_properties.modified=doc.core_properties.created
    doc.core_properties.subject="Review draft; genuine human validation and investigator sign-off pending"
    output.parent.mkdir(parents=True,exist_ok=True); doc.save(output)
    # Structural checks, separate from visual QA.
    check=Document(output)
    assert len(check.tables)==len(doc.tables)
    for table in check.tables:
        grid=[int(c.get(qn("w:w"))) for c in table._tbl.tblGrid.gridCol_lst]
        assert sum(grid)==9360
        for row in table.rows:
            assert [int(c._tc.tcPr.tcW.get(qn("w:w"))) for c in row.cells]==grid
    assert "{{" not in "\n".join(p.text for p in check.paragraphs)
    report={"source_sha256":hashlib.sha256(source.read_bytes()).hexdigest(),
            "docx_sha256":hashlib.sha256(output.read_bytes()).hexdigest(),"tables":len(check.tables),
            "preset":"narrative_proposal","header_pattern":"memo_masthead",
            "named_overrides":["academic_black_hierarchy","compact_22pt_title","9pt_table_text","no_rules"],
            "structural_checks_passed":True,"visual_qa_completed":False}
    (output.parent / "document_build.json").write_text(json.dumps(report,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(report,indent=2))


if __name__=="__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source",type=Path); parser.add_argument("output",type=Path)
    args=parser.parse_args(); build(args.source,args.output)
